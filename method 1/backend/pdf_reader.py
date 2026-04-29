import os
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation

def read_pdf(file_path):
    text = ""
    reader = PdfReader(file_path)
    for page in reader.pages:
        text += page.extract_text() or ""
    return text

def read_docx(file_path):
    text = ""
    doc = Document(file_path)
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def read_ppt(file_path):
    text = ""
    prs = Presentation(file_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def read_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return read_pdf(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    elif ext == ".pptx":
        return read_ppt(file_path)
    elif ext == ".txt":
        return read_txt(file_path)
    else:
        return "Unsupported file format"